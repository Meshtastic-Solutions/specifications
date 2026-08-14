#!/usr/bin/env python3
"""Regenerate the markdown rendering of the device-behavior audit spreadsheet.

The .xlsx is the authoritative artifact; README.md, develop.md and master.md in this
directory are derived from it. Re-run this after updating the spreadsheet:

    pip install openpyxl
    python3 audit/convert_xlsx_to_md.py

Run from the repository root, or set SRC/OUT below.
"""
import openpyxl, re, os, collections

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'meshtastic_device_behavior_audit.xlsx')
OUT = HERE
wb = openpyxl.load_workbook(SRC)

HDR = [c.value for c in wb['Develop'][1]]
COL = {h: i for i, h in enumerate(HDR)}

# Columns shown in the per-branch overview table, as (header, display name).
OVERVIEW = [
    ('User button (BUTTON_PIN)', 'User button'),
    ('Alt/Back button', 'Alt/back'),
    ('Cancel button', 'Cancel'),
    ('Touch button', 'Touch'),
    ('Power/status LED', 'Status LED'),
    ('LED polarity (LED_STATE_ON)', 'Polarity'),
    ('Notification LED', 'Notif LED'),
    ('Buzzer', 'Buzzer'),
    ('Touchscreen', 'Screen'),
]


def slug(text):
    s = text.lower()
    s = re.sub(r'[^a-z0-9 _-]', '', s)
    return s.replace(' ', '-')


def esc(text):
    """Escape a value for use inside a markdown table cell."""
    return str(text).replace('|', r'\|').replace('\n', '<br>')


def short(val, limit=30):
    """Compact a define down to its resolved GPIO (or symbol) for the overview table.

    Values arrive in several shapes: 'BUTTON_PIN=38', 'BUTTON_PIN=PIN_BUTTON1((32 + 10) = 42)',
    'LED_POWER=PIN_LED1(4)', 'PIN_BUZZER=(32 + 10) = 42', 'HAS_PMU=derived (AXP192)'. Where the
    audit resolved an arithmetic or symbolic pin to a number, show that number; otherwise show
    the symbol verbatim (never digits scraped out of a part name like AXP192 or PA8).
    """
    if not val:
        return ''
    lines = [x.strip() for x in str(val).split('\n') if x.strip()]
    first = lines[0]
    if '=' in first:
        first = first.split('=', 1)[1].strip()
    out = None
    if re.fullmatch(r'-?\d+', first):                       # 38 / -1
        out = first
    elif (m := re.findall(r'=\s*(-?\d+)\s*\)*$', first)):   # ... = 42 / (= 42) / (D18 (= 18))
        out = m[-1]
    elif (m := re.fullmatch(r'[A-Za-z_]\w*\((\d+)\)', first)):  # PIN_LED1(4)
        out = m.group(1)
    if out is None:
        out = first if len(first) <= limit else first[:limit - 1] + '…'
    # flag that this variant declares more than one define in the field
    return out + (' +' if len(lines) > 1 else '')


def rows(sheet):
    ws = wb[sheet]
    out = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c + 1).value for c in range(len(HDR))]
        if not vals[0]:
            continue
        out.append({h: (str(v).strip() if v is not None else '') for h, v in zip(HDR, vals)})
    return out


def marker(rec):
    vs = rec['Vs other branch']
    if vs.startswith('ONLY ON'):
        return '🟢'
    if vs.startswith('DIFFERS'):
        return '🟡'
    return ''


def behavior_bullets(text):
    """The behavior column is a '•'-bulleted list with wrapped continuation lines."""
    bullets = []
    for line in text.split('\n'):
        line = line.strip()
        if not line:
            continue
        if line.startswith('•'):
            bullets.append(line.lstrip('•').strip())
        elif bullets:
            bullets[-1] += ' ' + line
        else:
            bullets.append(line)
    return bullets


def branch_file(sheet, other, path):
    recs = rows(sheet)
    L = []
    L.append(f'# Device behavior audit — `{sheet.lower()}` branch\n')
    L.append(f'One entry per variant: **{len(recs)} variants** across '
             f'{len(set(r["Platform"] for r in recs))} platform families, with the effective '
             f'(post-include-chain) defines, the behavior those defines produce in the consuming '
             f'code, and the raw defines they were resolved from.\n')
    L.append(f'Generated from `meshtastic_device_behavior_audit.xlsx` (sheet “{sheet}”). '
             f'See [the audit overview](README.md) for method, semantics, and caveats.\n')
    L.append(f'🟢 = variant exists only on `{sheet.lower()}` · '
             f'🟡 = defines differ from `{other.lower()}`\n')

    # ---- platform breakdown
    plat = collections.Counter(r['Platform'] for r in recs)
    L.append('## Platforms\n')
    L.append('| Platform | Variants |')
    L.append('|---|---:|')
    for p, n in sorted(plat.items(), key=lambda kv: (-kv[1], kv[0])):
        L.append(f'| `{p}` | {n} |')
    L.append('')

    # ---- overview table
    L.append('## Overview\n')
    L.append('A condensed index — pin numbers are the resolved GPIO, blank means the variant does '
             'not declare that feature, and `+` means it declares more in that field than fits '
             'here. Variant names link to the full entry below, which is authoritative.\n')
    L.append('| Variant | ' + ' | '.join(d for _, d in OVERVIEW) + ' |')
    L.append('|---|' + '---|' * len(OVERVIEW))
    for rec in recs:
        name = rec['Variant']
        cells = [short(rec[h]) for h, _ in OVERVIEW]
        cells = [f'`{esc(c)}`' if c else '' for c in cells]
        mark = marker(rec)
        label = f'[`{name}`](#{slug(name)})' + (f' {mark}' if mark else '')
        L.append(f'| {label} | ' + ' | '.join(cells) + ' |')
    L.append('')

    # ---- per-variant detail
    L.append('## Variants\n')
    for rec in recs:
        name = rec['Variant']
        mark = marker(rec)
        L.append(f'### {name}\n')
        meta = [f'Platform `{rec["Platform"]}`']
        if rec['Vs other branch']:
            meta.append(f'{mark} **{esc(rec["Vs other branch"])}**')
        L.append(' · '.join(meta) + '\n')

        declared = [(h, rec[h]) for h in HDR
                    if h not in ('Variant', 'Platform', 'Behavior (resolved from consuming code)',
                                 'All matched defines (raw)', 'Vs other branch') and rec[h]]
        if declared:
            L.append('| Declared | Effective value |')
            L.append('|---|---|')
            for h, v in declared:
                lines = [f'`{esc(x.strip())}`' for x in str(v).split('\n') if x.strip()]
                L.append(f'| {h} | ' + '<br>'.join(lines) + ' |')
            L.append('')

        bullets = behavior_bullets(rec['Behavior (resolved from consuming code)'])
        if bullets:
            L.append('**Behavior**\n')
            for b in bullets:
                L.append(f'- {b}')
            L.append('')

        raw = rec['All matched defines (raw)']
        if raw:
            L.append('<details><summary>All matched defines (raw)</summary>\n')
            L.append('```c')
            L.extend(x.rstrip() for x in raw.split('\n'))
            L.append('```\n')
            L.append('</details>\n')

    with open(path, 'w') as f:
        f.write('\n'.join(L))
    return len(recs)


def hub(path, n_dev, n_mas):
    ws = wb['Branch Diff']
    diff = [[('' if v is None else str(v)) for v in row]
            for row in ws.iter_rows(values_only=True)][1:]
    diff = [d for d in diff if d[0]]

    L = []
    L.append('# Meshtastic device-behavior audit\n')
    L.append('Variant defines and their resolved behavior — the per-variant evidence base for the '
             'specification series. This is a markdown rendering of '
             '[`meshtastic_device_behavior_audit.xlsx`](../meshtastic_device_behavior_audit.xlsx), '
             'readable directly on GitHub; the spreadsheet remains the authoritative artifact and '
             'is the one to use for sorting and filtering.\n')
    L.append('| Sheet | Rendered as | Rows |')
    L.append('|---|---|---:|')
    L.append(f'| Develop | [`develop.md`](develop.md) | {n_dev} variants |')
    L.append(f'| Master | [`master.md`](master.md) | {n_mas} variants |')
    L.append(f'| Branch Diff | [below](#branch-diff) | {len(diff)} changes |')
    L.append('| Read Me | this page | — |')
    L.append('')
    L.append('For the narrative analysis built on this data — findings, branch drift, and the '
             'behavioral write-up — see '
             '[`meshtastic_device_behavior_audit_report.md`](../meshtastic_device_behavior_audit_report.md).\n')

    L.append('## Source and scope\n')
    L.append('**Source of truth:** `meshtastic/firmware` source code, branches `master` and '
             '`develop` (snapshots uploaded 2026-07-13).\n')
    L.append('**Scope:** LED, button, buzzer, vibration, NeoPixel/RGB, and touch-related defines in '
             '`variants/*/*/variant.h`, locally included shared headers (e.g. `station_common.h`), '
             "and `-D` build flags in each board's `platformio.ini`.\n")

    L.append("## How “effective” values were derived\n")
    L.append("Mirrors the firmware's include chain:\n")
    L.append('1. Raw defines parsed from `variant.h` + `platformio.ini`. First definition wins; '
             '`#if` context recorded.')
    L.append('2. nRF52 `architecture.h` mappings applied: `PIN_LED1`→`LED_POWER` (if unset), '
             '`PIN_BUTTON1`→`BUTTON_PIN`, `PIN_BUTTON_TOUCH`→`BUTTON_PIN_TOUCH`.')
    L.append('3. `configuration.h` defaults applied: `PIN_BUTTON2`→`ALT_BUTTON_PIN` (unless '
             '`CANCEL_BUTTON_PIN`), `LED_STATE_ON` defaults to 1 (active-high), '
             '`BUTTON_ACTIVE_LOW`/`PULLUP` default true.')
    L.append("4. Symbolic pin values resolved via the variant's `pins_arduino.h` where present; "
             "nRF port pins shown as arithmetic with the evaluated GPIO number "
             "(e.g. `(32 + 10) = 42` is P1.10).\n")

    L.append('## Behavior semantics\n')
    L.append('From consuming code, identical on both branches unless noted in the report.\n')
    L.append('| Feature | Behavior |')
    L.append('|---|---|')
    L.append('| **User button** (`src/input/InputBroker.cpp` + `ButtonThread`) | With screen: '
             'click=`USER_PRESS` (cycle UI), 500 ms hold=`SELECT`, 3.9 s hold=`SHUTDOWN`. Without '
             'screen: click=`USER_PRESS`, double=`SEND_PING`, triple=`GPS_TOGGLE`, '
             '3.9 s=`SHUTDOWN` (ThinkNode M7: +15 s=`FACTORY_RESET`). |')
    L.append('| **Alt/back button** | click=`ALT_PRESS`, 500 ms=`ALT_LONG`. |')
    L.append('| **Cancel button** | click=`CANCEL`, 4 s=`SHUTDOWN`. |')
    L.append('| **Touch button** | 500 ms=`BACK`; on boards with `PIN_EINK_EN` it instead '
             'momentarily lights the e-ink backlight while held. |')
    L.append('| **Power/status LED** (`src/modules/StatusLEDModule.cpp`) | discharging=1 ms blip '
             'per 1 s heartbeat; charging=1 s toggle; charged=solid; battery <5%=2 s burst of '
             '250 ms blinks every 30 s. Disabled by the user config flag `led_heartbeat_disabled`. |')
    L.append('| **Pairing LED** | BLE unpaired=2 s blink, pairing=1 s blink, connected=solid; '
             'turns off 30 s after the last state change. |')
    L.append('| **Notification LED / buzzer / vibration** | Default `ExternalNotificationModule` '
             'outputs (`NodeDB.cpp` defaults: LED 1000 ms on message, buzzer PWM ringtone, '
             'vibration 500 ms with 2 s nag). |')
    L.append('| **LoRa RX LED** | 100 ms flash per received packet. |')
    L.append('| **WiFi/BLE LEDs** | On while connected. |')
    L.append('| **InkHUD variants** | When the InkHUD display mode is active, button handling moves '
             "to the niche `TwoButton` driver wired in the variant's `nicheGraphics.h` (see the "
             'behavior entry for each variant). |')
    L.append('')

    L.append('## Caveats\n')
    L.append('- Defines inside `#if` blocks reflect the base build; shield/sub-build overrides '
             '(e.g. T-Beam + ST7796 display shield) are recorded with their condition in the raw '
             'defines but not resolved.')
    L.append("- Symbols provided by the Arduino core's own `pins_arduino.h` (board packages not "
             "vendored in the repo, e.g. heltec_v3's `LED`) stay symbolic.")
    L.append("- “Candidate dead defines” are LED symbols (`LED_BLUE`/`GREEN`/`RED`, `USER_LED`) "
             'defined by a variant but never read by Meshtastic application code on either branch '
             '(they may still matter to the bootloader or Arduino core).\n')

    L.append('## Branch diff\n')
    L.append('Every added variant and changed define between `master` and `develop`.\n')
    L.append('| Variant | Change | Define | master value | develop value |')
    L.append('|---|---|---|---|---|')
    for d in diff:
        d = (d + [''] * 5)[:5]
        cells = [f'`{esc(d[0])}`', esc(d[1])]
        cells += [f'`{esc(x)}`' if x else '' for x in d[2:5]]
        L.append('| ' + ' | '.join(cells) + ' |')
    L.append('')
    L.append('In the per-branch files these variants are flagged 🟢 (exists only on that branch) '
             'or 🟡 (defines differ between branches).\n')

    with open(path, 'w') as f:
        f.write('\n'.join(L))


os.makedirs(OUT, exist_ok=True)
n_dev = branch_file('Develop', 'Master', f'{OUT}/develop.md')
n_mas = branch_file('Master', 'Develop', f'{OUT}/master.md')
hub(f'{OUT}/README.md', n_dev, n_mas)
for f in ('README.md', 'develop.md', 'master.md'):
    print(f'{f:12s} {os.path.getsize(OUT + "/" + f)/1024:8.1f} KB')
